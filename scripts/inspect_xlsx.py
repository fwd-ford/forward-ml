"""Quick structural inspection of the official Ford xlsx files.
Goal: schema, row count, sheet inventory, sample rows. No analytics yet."""
import sys
from openpyxl import load_workbook

def inspect(path: str) -> None:
    print(f"\n===== {path} =====")
    wb = load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n--- sheet: {name} ---")
        print(f"max_row={ws.max_row}  max_col={ws.max_column}")
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
            print(f"header ({len([h for h in header if h is not None])} non-empty):")
            for i, h in enumerate(header):
                print(f"  [{i:>2}] {h!r}")
            print("first 3 data rows:")
            for k, r in enumerate(rows):
                if k >= 3:
                    break
                preview = [str(v)[:40] if v is not None else None for v in r]
                print(f"  {preview}")
        except StopIteration:
            print("(empty)")
    wb.close()

if __name__ == "__main__":
    for p in sys.argv[1:]:
        inspect(p)
